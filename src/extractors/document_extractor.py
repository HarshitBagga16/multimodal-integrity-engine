
from pathlib import Path

from src.extractors.video_extractor import extract_video
import pymupdf
from docx import Document


def extract_pdf(file_path: Path) -> list[dict]:
    """Extract non-empty PDF pages with source metadata."""
    extracted = []

    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()

            if text:
                extracted.append({
                    "text": text,
                    "source_file": file_path.name,
                    "file_type": "pdf",
                    "page_number": page_number,
                    "timestamp": None,
                })

    return extracted


def extract_docx(file_path: Path) -> list[dict]:
    """Extract non-empty DOCX paragraphs and table rows."""
    document = Document(file_path)
    extracted = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted.append({
                "text": text,
                "source_file": file_path.name,
                "file_type": "docx",
                "page_number": None,
                "timestamp": None,
            })

    # Word tables may contain meaningful submission content.
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if text:
                extracted.append({
                    "text": text,
                    "source_file": file_path.name,
                    "file_type": "docx",
                    "page_number": None,
                    "timestamp": None,
                })

    return extracted

def extract_document(file_path: Path) -> list[dict]:
    """Dispatch to the correct extractor."""
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".mp4":
        return extract_video(file_path)

    raise ValueError(f"Unsupported document format: {extension}")